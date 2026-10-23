#!/usr/bin/env python3
"""Manual, read-only classification of a Banksalad transaction export.

The workbook is evidence, not an accounting ledger. No row is discarded and
classification never changes the signed value supplied by Banksalad.

Install the Python dependency with ``python3 -m pip install -r
scripts/tools/banksalad-requirements.txt`` before running this tool or its tests.
"""

import argparse
import calendar
from collections import Counter, defaultdict
from datetime import date, datetime, time, timedelta
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import re

from openpyxl import load_workbook

SHEET_NAME = "가계부 내역"
HEADERS = ("날짜", "시간", "타입", "대분류", "소분류", "내용", "금액", "화폐", "결제수단", "메모")
CLASSIFICATIONS = (
    "income", "zero-value informational income", "expense", "card settlement", "own-account transfer",
    "verified-family transfer", "owner-estimated family transfer", "account verification credit",
    "savings contribution", "investment contribution",
    "debt-service-unsplit", "refund/reversal", "third-party purchase pass-through", "review",
    "expense reimbursement received",
    "wallet top-up for purchase", "wallet top-up", "unverified-destination transfer",
)
RAW_TRANSACTION_FIELDS = ("date", "time", "original_type", "original_category",
                          "original_subcategory", "description", "signed_amount",
                          "currency", "payment_instrument", "memo")
APPROVED_SOURCE_SHA256 = "698322936b32008f401242c68043b4fe6f1119c2d0da585f04a94d7a552d59b2"
SOURCE_COVERAGE_BY_SHA256 = {
    APPROVED_SOURCE_SHA256: ("2025-10-21", "2026-10-21"),
}
EXPORT_DATE_RANGE = re.compile(r"(\d{4}-\d{2}-\d{2})_(\d{4}-\d{2}-\d{2})\Z")
APPROVED_DUPLICATE_DATE = "2026-10-22"
APPROVED_DUPLICATE_PAIRS = ((101, 103), (269, 271), (840, 842), (861, 862),
                            (1018, 1020), (1338, 1340), (1434, 1436), (1720, 1722))
APPROVED_UNVERIFIED_KAKAO_ROWS = (247, 348, 505, 786, 1451)
APPROVED_UNVERIFIED_KAKAO_TOTAL = Decimal("400000")
ONE_OFF_TAX_MEMOS = {"주민세", "재산세", "지방소득세", "2024 연말정산"}


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def amount_value(value):
    if isinstance(value, bool) or value is None:
        return None
    try:
        amount = Decimal(str(value).replace(",", ""))
    except InvalidOperation:
        return None
    return amount if amount.is_finite() else None


def amount_text(amount):
    return format(amount, "f") if amount is not None else ""


def classify(transaction):
    original_type = transaction["original_type"]
    category = transaction["original_category"]
    subcategory = transaction["original_subcategory"]
    original_description = transaction["description"]
    description = original_description.strip()
    memo = transaction["memo"]
    amount = amount_value(transaction["signed_amount"])
    if (amount == 0 and original_type == "수입" and category == "금융수입"
            and original_description == "입출금통장 이자"):
        return "zero-value informational income", ["zero_amount", "named_bank_interest_information"]
    if amount is None or amount == 0:
        return "review", ["missing_or_zero_amount"]

    if original_type == "이체" and amount < 0 and original_description == "인천이음카드":
        return "own-account transfer", ["named_regional_currency_top_up", "outflow"]

    # Bank export category identifies the settlement; counting it as a purchase
    # would double-count spending already recorded on the card instrument.
    if category == "카드대금":
        if original_type == "이체" and amount < 0:
            return "card settlement", ["banksalad_card_settlement_category", "outflow"]
        return "review", ["card_settlement_category_with_unexpected_type_or_direction"]
    if original_type == "지출" and amount > 0:
        if category in {"이체", "내계좌이체", "저축"} or subcategory == "이자/대출":
            return "review", ["positive_expense_direction", "contradictory_transfer_or_settlement_label"]
        return "refund/reversal", ["banksalad_expense_type", "positive_amount_reversal_direction"]
    if subcategory == "이자/대출" or "대출원리금상환" in description:
        return "debt-service-unsplit", ["loan_payment_label", "principal_interest_unknown"]
    if "송금 취소" in description or "환불" in description or "결제취소" in description:
        return "refund/reversal", ["description_indicates_reversal", "direction_unverified"]
    if amount < 0 and description == "주택청약":
        return "savings contribution", ["named_housing_savings_product", "outflow"]
    if amount < 0 and amount == -250000 and re.fullmatch(r"한국증권\d{4}", description):
        return "investment contribution", ["recurring_korea_securities_250000", "outflow"]
    if amount < 0 and description == "업비트 입금":
        return "investment contribution", ["named_upbit_deposit", "outflow", "destination_unverified"]

    # A memo may mention family, but it is not proof of who held the counterparty
    # account. Stage 1 has no verified family identity roster or account mapping.
    if category == "내계좌이체" and original_type == "이체":
        return "review", ["banksalad_own_account_category", "opposite_account_leg_unverified"]
    if original_type == "이체" or category in {"이체", "저축"}:
        return "review", ["transfer_purpose_or_counterparty_unverified"]
    if category == "미분류" or subcategory == "미분류" and original_type not in {"수입", "지출"}:
        return "review", ["unclassified_source_category"]
    if original_type == "수입" and amount > 0:
        return "income", ["banksalad_income_type", "positive_amount"]
    if original_type == "지출" and amount < 0:
        return "expense", ["banksalad_expense_type", "negative_amount"]
    return "review", ["type_amount_conflict_or_unknown"]


def load_transactions(workbook_path):
    source_hash = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if SHEET_NAME not in workbook:
            raise ValueError(f"Required sheet missing: {SHEET_NAME}")
        sheet = workbook[SHEET_NAME]
        rows = sheet.iter_rows(values_only=True)
        header = tuple(cell_text(value).strip() for value in next(rows))
        if header[:len(HEADERS)] != HEADERS:
            raise ValueError("Unexpected transaction columns")
        transactions = []
        for row_number, values in enumerate(rows, start=2):
            fields = dict(zip(HEADERS, values))
            if all(value is None for value in values[:len(HEADERS)]):
                continue
            amount = amount_value(fields["금액"])
            transaction = {
                "source_sheet": SHEET_NAME,
                "source_row": row_number,
                "source_file_sha256": source_hash,
                "date": cell_text(fields["날짜"])[:10],
                "time": cell_text(fields["시간"]),
                "original_type": cell_text(fields["타입"]),
                "original_category": cell_text(fields["대분류"]),
                "original_subcategory": cell_text(fields["소분류"]),
                "description": cell_text(fields["내용"]),
                "signed_amount": amount_text(amount) if amount is not None else cell_text(fields["금액"]),
                "currency": cell_text(fields["화폐"]),
                "payment_instrument": cell_text(fields["결제수단"]),
                "memo": cell_text(fields["메모"]),
            }
            transaction["classification"], transaction["classification_evidence"] = classify(transaction)
            transaction["exact_duplicate_rows"] = []
            transaction["aggregation_status"] = "included"
            transactions.append(transaction)
        return source_hash, transactions
    finally:
        workbook.close()


def mark_exact_duplicates(transactions):
    groups = defaultdict(list)
    for transaction in transactions:
        key = tuple(transaction[field] for field in RAW_TRANSACTION_FIELDS)
        groups[key].append(transaction)
    duplicate_groups = []
    for group in groups.values():
        if len(group) < 2:
            continue
        source_rows = [transaction["source_row"] for transaction in group]
        duplicate_groups.append({"source_rows": source_rows, "count": len(group),
                                 "midnight_transit": all(
                                     item["time"].startswith("00:00")
                                     and item["original_category"] == "교통"
                                     for item in group)})
        for transaction in group:
            transaction["exact_duplicate_rows"] = [row for row in source_rows
                                                   if row != transaction["source_row"]]
    return sorted(duplicate_groups, key=lambda group: group["source_rows"][0])


def mark_paired_own_account_transfers(transactions):
    """Pair bank-side legs; unverified wallet top-ups remain in review.

    The approved 인천이음카드 top-up is classified earlier by classify().
    """
    groups = defaultdict(list)
    wallet_terms = ("PAYCO", "카카오페이", "포인트", "충전", "업비트")
    for transaction in transactions:
        if transaction["original_category"] != "내계좌이체" or transaction["classification"] != "review":
            continue
        description_and_instrument = transaction["description"] + transaction["payment_instrument"]
        if any(term in description_and_instrument for term in wallet_terms):
            continue
        amount = amount_value(transaction["signed_amount"])
        if amount is not None:
            groups[(transaction["date"], abs(amount))].append(transaction)
    for group in groups.values():
        if len(group) != 2:
            continue
        first, second = group
        if (amount_value(first["signed_amount"]) != -amount_value(second["signed_amount"])
                or first["payment_instrument"] == second["payment_instrument"]
                or not first["payment_instrument"] or not second["payment_instrument"]):
            continue
        for transaction in group:
            transaction["classification"] = "own-account transfer"
            transaction["classification_evidence"] = [
                "banksalad_own_account_category", "opposite_signed_same_day_bank_legs",
                "distinct_payment_instruments",
            ]


def transaction_time(transaction):
    try:
        return datetime.fromisoformat(f"{transaction['date']}T{transaction['time']}")
    except (TypeError, ValueError):
        raise ValueError(f"Invalid transaction timestamp for row {transaction['source_row']}") from None


def validate_payment_decisions(transactions, source_hash, decisions, reserved_rows):
    """Check the whole payment decision set before apply_owner_decisions changes any row."""
    by_row = {item["source_row"]: item for item in transactions}
    used_rows = set(reserved_rows)
    payment_rows = set()
    evidence_fields = {"source_row", "expected_date", "expected_time", "expected_description",
                       "expected_signed_amount", "expected_payment_instrument", "expected_memo",
                       "expected_original_review_classification"}

    approved_receipt_rows = set(decisions["approved_own_transfer_rows"])

    def checked_row(evidence, expected_classification, role, allow_approved_receipt=False):
        if not isinstance(evidence, dict) or set(evidence) != evidence_fields:
            raise ValueError(f"Invalid {role} evidence fields")
        row = evidence["source_row"]
        if (type(row) is not int or row < 2 or row in payment_rows
                or (row in used_rows and not (allow_approved_receipt and row in approved_receipt_rows))):
            raise ValueError(f"Duplicate or conflicting {role} row")
        used_rows.add(row)
        payment_rows.add(row)
        transaction = by_row.get(row)
        if (transaction is None or transaction["source_file_sha256"] != source_hash
                or evidence["expected_original_review_classification"] != expected_classification
                or transaction["classification"] != expected_classification
                or any(not isinstance(value, str) for key, value in evidence.items() if key != "source_row")
                or any(transaction[field] != evidence["expected_" + field]
                       for field in ("date", "time", "description", "signed_amount",
                                     "payment_instrument", "memo"))):
            raise ValueError(f"{role} row {row} evidence mismatch")
        return transaction

    top_ups = decisions.get("confirmed_kakao_payment_top_ups", [])
    generic = decisions.get("unverified_kakao_transfers", [])
    expenses = decisions.get("confirmed_payment_transfer_expenses", [])
    if any(not isinstance(group, list) for group in (top_ups, generic, expenses)):
        raise ValueError("Invalid payment decision group")
    top_up_ids = set()
    for group in top_ups:
        if (not isinstance(group, dict)
                or set(group) != {"group_id", "bank_outflow", "wallet_receipt", "purchase"}
                or not isinstance(group["group_id"], str) or not group["group_id"].strip()
                or group["group_id"] in top_up_ids):
            raise ValueError("Invalid Kakao payment top-up group")
        top_up_ids.add(group["group_id"])
        bank = checked_row(group["bank_outflow"], "review", "bank outflow")
        receipt = checked_row(group["wallet_receipt"], "review", "wallet receipt",
                              allow_approved_receipt=True)
        if receipt["source_row"] not in approved_receipt_rows:
            raise ValueError("Wallet receipt lacks prior own-transfer approval")
        purchase = checked_row(group["purchase"], "expense", "existing purchase")
        bank_amount = amount_value(bank["signed_amount"])
        receipt_amount = amount_value(receipt["signed_amount"])
        purchase_amount = amount_value(purchase["signed_amount"])
        if (bank["description"] != "카카오페이"
                or bank["payment_instrument"] != "하나은행 입출금통장"
                or receipt["payment_instrument"] != "카카오페이 머니"
                or purchase["payment_instrument"] != "카카오페이 머니"
                or bank_amount is None or bank_amount >= 0
                or receipt_amount != -bank_amount
                or purchase_amount is None or purchase_amount >= 0
                or abs(transaction_time(receipt) - transaction_time(bank)) > timedelta(seconds=60)
                or abs(transaction_time(purchase) - transaction_time(bank)) > timedelta(seconds=60)):
            raise ValueError(f"Kakao payment top-up companion mismatch: {group['group_id']}")
    for evidence in generic:
        bank = checked_row(evidence, "review", "unverified transfer")
        amount = amount_value(bank["signed_amount"])
        if (bank["description"] != "카카오페이"
                or bank["payment_instrument"] != "하나은행 입출금통장"
                or amount is None or amount >= 0):
            raise ValueError(f"Invalid unverified transfer row {bank['source_row']}")
        timestamp = transaction_time(bank)
        if any(item["source_row"] != bank["source_row"]
               and item["payment_instrument"] == "카카오페이 머니"
               and abs(transaction_time(item) - timestamp) <= timedelta(minutes=10)
               for item in transactions):
            raise ValueError(f"Wallet usage found near unverified transfer row {bank['source_row']}")
    if source_hash == APPROVED_SOURCE_SHA256:
        if (tuple(item["source_row"] for item in generic) != APPROVED_UNVERIFIED_KAKAO_ROWS
                or sum((-amount_value(by_row[item["source_row"]]["signed_amount"])
                        for item in generic), Decimal(0)) != APPROVED_UNVERIFIED_KAKAO_TOTAL):
            raise ValueError("Approved unverified Kakao expense rows or total mismatch")
    for evidence in expenses:
        bank = checked_row(evidence, "review", "payment expense")
        amount = amount_value(bank["signed_amount"])
        if (amount is None or amount >= 0
                or not ((bank["description"] == "카카오페이"
                         and bank["payment_instrument"] == "하나은행 입출금통장")
                        or (bank["description"] == "네이버페이"
                            and bank["payment_instrument"] == "페이코 간편결제"))):
            raise ValueError(f"Invalid payment expense row {bank['source_row']}")
        if bank["description"] == "카카오페이":
            timestamp = transaction_time(bank)
            if any(item["source_row"] != bank["source_row"]
                   and (amount_value(item["signed_amount"]) is not None
                        and amount_value(item["signed_amount"]) < 0)
                   and abs(transaction_time(item) - timestamp) <= timedelta(minutes=10)
                   for item in transactions):
                raise ValueError(f"Separate charge near payment expense row {bank['source_row']}")
    return top_ups, generic, expenses


def validate_duplicate_decisions(transactions, source_hash, decisions):
    """Only an owner-approved, source-bound exact pair can change aggregation."""
    if "confirmed_exact_duplicate_pairs" in decisions:
        pairs = decisions["confirmed_exact_duplicate_pairs"]
    elif source_hash == APPROVED_SOURCE_SHA256:
        pairs = [{"keep_source_row": keep, "exclude_source_row": exclude,
                  "approval_date": APPROVED_DUPLICATE_DATE, "treatment": "count_first_once"}
                 for keep, exclude in APPROVED_DUPLICATE_PAIRS]
    else:
        pairs = []
    if not isinstance(pairs, list):
        raise ValueError("Invalid exact duplicate decisions")
    if source_hash == APPROVED_SOURCE_SHA256 and [
            (pair.get("keep_source_row"), pair.get("exclude_source_row"))
            for pair in pairs if isinstance(pair, dict)] != list(APPROVED_DUPLICATE_PAIRS):
        raise ValueError("Approved exact duplicate source row ids mismatch")
    by_row = {item["source_row"]: item for item in transactions}
    used_rows = set()
    for pair in pairs:
        if (not isinstance(pair, dict)
                or set(pair) != {"keep_source_row", "exclude_source_row", "approval_date", "treatment"}):
            raise ValueError("Invalid exact duplicate decision fields")
        keep_row, exclude_row = pair["keep_source_row"], pair["exclude_source_row"]
        if (type(keep_row) is not int or type(exclude_row) is not int
                or keep_row < 2 or exclude_row <= keep_row
                or keep_row in used_rows or exclude_row in used_rows):
            raise ValueError("Invalid or reused exact duplicate source row ids")
        used_rows.update((keep_row, exclude_row))
        try:
            approval_date = date.fromisoformat(pair["approval_date"])
        except (TypeError, ValueError):
            raise ValueError("Invalid exact duplicate approval date") from None
        if (approval_date.isoformat() != pair["approval_date"]
                or pair["treatment"] != "count_first_once"
                or (source_hash == APPROVED_SOURCE_SHA256
                    and pair["approval_date"] != APPROVED_DUPLICATE_DATE)):
            raise ValueError("Invalid exact duplicate approval or treatment")
        keep, exclude = by_row.get(keep_row), by_row.get(exclude_row)
        if (keep is None or exclude is None
                or keep["source_file_sha256"] != source_hash
                or exclude["source_file_sha256"] != source_hash
                or keep["original_category"] == "교통"
                or exclude["original_category"] == "교통"
                or any(keep[field] != exclude[field] for field in RAW_TRANSACTION_FIELDS)):
            raise ValueError(f"Exact duplicate source evidence mismatch: {keep_row}, {exclude_row}")
    return pairs


def apply_owner_decisions(transactions, source_hash, decisions):
    """Validate all source-scoped owner decisions before annotating any row."""
    required_fields = {"source_file_sha256", "approval_date", "approved_own_transfer_rows"}
    if (not isinstance(decisions, dict) or not required_fields <= set(decisions)
            or set(decisions) - required_fields
            - {"confirmed_destinations", "confirmed_irp_contributions", "confirmed_review_decisions",
               "confirmed_account_verification_credits", "tentative_review_contexts",
               "confirmed_review_purposes", "owner_estimated_family_transfers",
               "confirmed_refund_reversals", "confirmed_consumption_expenses",
               "confirmed_third_party_purchase_pass_throughs", "confirmed_own_account_receipts",
               "confirmed_wallet_top_up_pairs", "confirmed_group_expense_reimbursements",
               "confirmed_kakao_payment_top_ups", "unverified_kakao_transfers",
               "confirmed_payment_transfer_expenses", "confirmed_bank_interest_income",
               "confirmed_prior_spending_reimbursements", "confirmed_food_consumption_expenses",
               "confirmed_unmatched_wallet_top_ups", "confirmed_exact_duplicate_pairs"}):
        raise ValueError("Invalid owner decisions fields")
    if decisions["source_file_sha256"] != source_hash:
        raise ValueError("Owner decisions source hash mismatch")
    try:
        approval_date = date.fromisoformat(decisions["approval_date"])
    except (TypeError, ValueError):
        raise ValueError("Invalid owner approval date") from None
    if approval_date.isoformat() != decisions["approval_date"]:
        raise ValueError("Invalid owner approval date")
    approved_rows = decisions["approved_own_transfer_rows"]
    if (not isinstance(approved_rows, list)
            or any(type(row) is not int or row < 2 for row in approved_rows)
            or len(approved_rows) != len(set(approved_rows))):
        raise ValueError("Invalid or duplicate approved row ids")
    confirmations = decisions.get("confirmed_destinations", [])
    if not isinstance(confirmations, list):
        raise ValueError("Invalid destination confirmations")
    confirmed_rows = set()
    for confirmation in confirmations:
        if (not isinstance(confirmation, dict)
                or set(confirmation) != {"source_row", "destination_account", "confirmation_date"}):
            raise ValueError("Invalid destination confirmation fields")
        row = confirmation["source_row"]
        if type(row) is not int or row in confirmed_rows or row not in approved_rows:
            raise ValueError("Destination row must uniquely identify an approved transfer")
        confirmed_rows.add(row)
        account = confirmation["destination_account"]
        if not isinstance(account, str) or not account.strip() or account != account.strip():
            raise ValueError("Invalid destination account")
        confirmation_date_text = confirmation["confirmation_date"]
        try:
            confirmation_date = date.fromisoformat(confirmation_date_text)
        except (TypeError, ValueError):
            raise ValueError("Invalid destination confirmation date") from None
        if confirmation_date.isoformat() != confirmation_date_text:
            raise ValueError("Invalid destination confirmation date")
    irp_confirmations = decisions.get("confirmed_irp_contributions", [])
    if not isinstance(irp_confirmations, list):
        raise ValueError("Invalid IRP contribution confirmations")
    by_row = {transaction["source_row"]: transaction for transaction in transactions}
    if len(by_row) != len(transactions):
        raise ValueError("Duplicate transaction source row ids")
    for row in approved_rows:
        transaction = by_row.get(row)
        if (transaction is None or transaction["classification"] != "review"
                or "banksalad_own_account_category" not in transaction["classification_evidence"]):
            raise ValueError(f"Approved row {row} is not pending own-account review")
    confirmed_irp_rows = set()
    for confirmation in irp_confirmations:
        if (not isinstance(confirmation, dict)
                or set(confirmation) != {"source_row", "destination_account",
                                         "contribution_purpose", "confirmation_date"}):
            raise ValueError("Invalid IRP contribution confirmation fields")
        row = confirmation["source_row"]
        if type(row) is not int or row < 2 or row in confirmed_irp_rows:
            raise ValueError("Invalid or duplicate IRP contribution row ids")
        confirmed_irp_rows.add(row)
        transaction = by_row.get(row)
        if (transaction is None or transaction["source_file_sha256"] != source_hash
                or transaction["classification"] != "investment contribution"
                or amount_value(transaction["signed_amount"]) != Decimal("-250000")
                or re.fullmatch(r"한국증권\d{4}", transaction["description"]) is None):
            raise ValueError(f"IRP row {row} is not the expected investment contribution")
        if (confirmation["destination_account"] != "한국투자증권 IRP"
                or confirmation["contribution_purpose"] != "IRP 납입"):
            raise ValueError(f"Invalid IRP destination or purpose for row {row}")
        confirmation_date_text = confirmation["confirmation_date"]
        try:
            confirmation_date = date.fromisoformat(confirmation_date_text)
        except (TypeError, ValueError):
            raise ValueError("Invalid IRP confirmation date") from None
        if confirmation_date.isoformat() != confirmation_date_text:
            raise ValueError("Invalid IRP confirmation date")
    review_decisions = decisions.get("confirmed_review_decisions", [])
    if not isinstance(review_decisions, list):
        raise ValueError("Invalid review decisions")
    review_rows = set()
    for decision in review_decisions:
        fields = {"source_row", "expected_description", "expected_signed_amount",
                  "expected_payment_instrument", "classification"}
        if (not isinstance(decision, dict)
                or set(decision) not in (fields, fields | {"purpose"})):
            raise ValueError("Invalid review decision fields")
        row = decision["source_row"]
        if type(row) is not int or row < 2 or row in review_rows or row in approved_rows:
            raise ValueError("Invalid or duplicate review decision row ids")
        review_rows.add(row)
        classification = decision["classification"]
        if (not isinstance(classification, str)
                or classification not in {"own-account transfer", "verified-family transfer", "expense"}):
            raise ValueError(f"Invalid review classification for row {row}")
        if ((classification == "expense" and decision.get("purpose") != "구독료")
                or (classification != "expense" and "purpose" in decision)):
            raise ValueError(f"Invalid review purpose for row {row}")
        transaction = by_row.get(row)
        amount = amount_value(transaction["signed_amount"]) if transaction is not None else None
        if (transaction is None or transaction["source_file_sha256"] != source_hash
                or transaction["classification"] != "review"
                or not all(isinstance(decision[field], str) for field in fields - {"source_row"})
                or transaction["description"] != decision["expected_description"]
                or transaction["signed_amount"] != decision["expected_signed_amount"]
                or transaction["payment_instrument"] != decision["expected_payment_instrument"]
                or amount is None or amount == 0
                or (classification == "expense" and amount >= 0)
                or (classification == "verified-family transfer" and amount <= 0)):
            raise ValueError(f"Review row {row} does not match owner-confirmed evidence")
    evidence_fields = {"source_row", "expected_date", "expected_description",
                       "expected_signed_amount", "expected_payment_instrument"}
    new_decision_groups = (
        ("confirmed_account_verification_credits", {"purpose"}),
        ("tentative_review_contexts", {"expected_memo", "owner_context"}),
        ("confirmed_review_purposes", {"purpose"}),
        ("owner_estimated_family_transfers", {"expected_memo", "owner_recollection",
                                              "candidate_related_row", "relation_status"}),
        ("confirmed_refund_reversals", {"expected_memo", "purpose", "spending_treatment"}),
        ("confirmed_consumption_expenses", {"expected_memo", "classification",
                                            "spending_treatment"}),
        ("confirmed_own_account_receipts", {"expected_memo",
                                            "expected_original_review_classification"}),
        ("confirmed_group_expense_reimbursements", {"expected_memo",
                                                    "expected_original_review_classification",
                                                    "owner_context", "spending_treatment"}),
        ("confirmed_bank_interest_income", {"expected_time", "expected_original_type",
                                             "expected_original_category", "expected_original_subcategory",
                                             "expected_currency", "expected_memo",
                                             "expected_original_review_classification"}),
        ("confirmed_prior_spending_reimbursements", {"expected_time", "expected_original_type",
                                                      "expected_original_category", "expected_original_subcategory",
                                                      "expected_currency", "expected_memo",
                                                      "expected_original_review_classification",
                                                      "spending_treatment"}),
        ("confirmed_food_consumption_expenses", {"expected_time", "expected_original_type",
                                                  "expected_original_category", "expected_original_subcategory",
                                                  "expected_currency", "expected_memo",
                                                  "expected_original_review_classification"}),
        ("confirmed_unmatched_wallet_top_ups", {"expected_time", "expected_original_type",
                                                  "expected_original_category", "expected_original_subcategory",
                                                  "expected_currency", "expected_memo",
                                                  "expected_original_review_classification"}),
    )
    new_decisions = {}
    new_rows = set()
    for group_name, additional_fields in new_decision_groups:
        group = decisions.get(group_name, [])
        if not isinstance(group, list):
            raise ValueError(f"Invalid {group_name}")
        new_decisions[group_name] = group
        for decision in group:
            if not isinstance(decision, dict) or set(decision) != evidence_fields | additional_fields:
                raise ValueError(f"Invalid {group_name} fields")
            row = decision["source_row"]
            if (type(row) is not int or row < 2 or row in new_rows
                    or row in approved_rows or row in review_rows):
                raise ValueError(f"Invalid or duplicate {group_name} row ids")
            new_rows.add(row)
            transaction = by_row.get(row)
            if (transaction is None or transaction["source_file_sha256"] != source_hash
                    or transaction["classification"] != "review"
                    or any(not isinstance(decision[field], str) for field in
                           (evidence_fields | additional_fields)
                           - {"source_row", "candidate_related_row"})
                    or transaction["date"] != decision["expected_date"]
                    or transaction["description"] != decision["expected_description"]
                    or transaction["signed_amount"] != decision["expected_signed_amount"]
                    or transaction["payment_instrument"] != decision["expected_payment_instrument"]
                    or ("expected_memo" in decision
                        and transaction["memo"] != decision["expected_memo"])):
                raise ValueError(f"Review row {row} does not match owner-confirmed evidence")
            amount = amount_value(transaction["signed_amount"])
            if (group_name == "confirmed_account_verification_credits"
                    and (amount != Decimal(1) or decision["purpose"] != "계좌개설·인증 입금")):
                raise ValueError(f"Invalid account verification credit for row {row}")
            if (group_name == "tentative_review_contexts"
                    and (not decision["owner_context"].strip()
                         or decision["owner_context"] != decision["owner_context"].strip())):
                raise ValueError(f"Invalid tentative owner context for row {row}")
            if (group_name == "confirmed_review_purposes"
                    and decision["purpose"] != "캐시백"):
                raise ValueError(f"Invalid confirmed review purpose for row {row}")
            if group_name == "owner_estimated_family_transfers":
                if (amount is None or amount == 0
                        or type(decision["candidate_related_row"]) is not int
                        or decision["candidate_related_row"] == row
                        or decision["relation_status"] != "candidate_only_unreconciled"
                        or not decision["owner_recollection"].strip()
                        or decision["owner_recollection"] != decision["owner_recollection"].strip()):
                    raise ValueError(f"Invalid owner-estimated family transfer for row {row}")
            if (group_name == "confirmed_refund_reversals"
                    and (amount is None or amount <= 0 or decision["purpose"] != "캐시백"
                         or decision["spending_treatment"] != "expense_offset")):
                raise ValueError(f"Invalid refund/reversal decision for row {row}")
            if (group_name == "confirmed_consumption_expenses"
                    and (amount is None or amount >= 0 or decision["classification"] != "expense"
                         or decision["spending_treatment"] != "consumption")):
                raise ValueError(f"Invalid consumption expense decision for row {row}")
            if group_name in {"confirmed_own_account_receipts",
                              "confirmed_group_expense_reimbursements"}:
                if (amount is None or amount <= 0
                        or decision["expected_original_review_classification"] != "review"
                        or transaction["classification"] != decision["expected_original_review_classification"]):
                    raise ValueError(f"Invalid confirmed receipt for row {row}")
            if (group_name == "confirmed_group_expense_reimbursements"
                    and (decision["owner_context"] != "group-expense settlement"
                         or decision["spending_treatment"] != "offset_pending_purchase_match")):
                raise ValueError(f"Invalid group expense reimbursement for row {row}")
            if group_name in {"confirmed_bank_interest_income",
                              "confirmed_prior_spending_reimbursements",
                              "confirmed_food_consumption_expenses",
                              "confirmed_unmatched_wallet_top_ups"}:
                source_fields = ("time", "original_type", "original_category",
                                 "original_subcategory", "currency", "memo")
                if (decision["expected_original_review_classification"] != "review"
                        or any(transaction[field] != decision[f"expected_{field}"]
                               for field in source_fields)
                        or (group_name in {"confirmed_bank_interest_income",
                                           "confirmed_prior_spending_reimbursements"}
                            and (amount is None or amount <= 0))
                        or (group_name in {"confirmed_food_consumption_expenses",
                                           "confirmed_unmatched_wallet_top_ups"}
                            and (amount is None or amount >= 0))
                        or (group_name == "confirmed_bank_interest_income"
                            and (transaction["description"] != "입금 내역"
                                 or transaction["payment_instrument"] != "하나은행 입출금통장"))
                        or (group_name == "confirmed_food_consumption_expenses"
                            and transaction["original_category"] != "식비")
                        or (group_name == "confirmed_prior_spending_reimbursements"
                            and decision["spending_treatment"] != "offset_pending_purchase_match")):
                    raise ValueError(f"Invalid final owner decision for row {row}")
    family_decisions = {item["source_row"]: item
                        for item in new_decisions["owner_estimated_family_transfers"]}
    for row, decision in family_decisions.items():
        other = family_decisions.get(decision["candidate_related_row"])
        if other is None or other["candidate_related_row"] != row:
            raise ValueError(f"Family transfer row {row} lacks reciprocal candidate context")
    pass_through_groups = decisions.get("confirmed_third_party_purchase_pass_throughs", [])
    if not isinstance(pass_through_groups, list):
        raise ValueError("Invalid third-party purchase pass-through groups")
    pass_through_rows = set()
    pass_through_ids = set()
    reserved_rows = set(approved_rows) | confirmed_irp_rows | review_rows | new_rows
    for group in pass_through_groups:
        if (not isinstance(group, dict)
                or set(group) != {"group_id", "expected_date", "expected_residual",
                                  "relation_status", "legs"}):
            raise ValueError("Invalid third-party purchase pass-through group fields")
        group_id = group["group_id"]
        if (not isinstance(group_id, str) or not group_id.strip()
                or group_id != group_id.strip() or group_id in pass_through_ids
                or group["relation_status"] != "owner_confirmed_purchase_reimbursement_not_exact_amount"
                or not isinstance(group["legs"], list) or len(group["legs"]) < 2):
            raise ValueError("Invalid third-party purchase pass-through group")
        pass_through_ids.add(group_id)
        try:
            group_date = date.fromisoformat(group["expected_date"])
        except (TypeError, ValueError):
            raise ValueError("Invalid third-party purchase pass-through date") from None
        if group_date.isoformat() != group["expected_date"]:
            raise ValueError("Invalid third-party purchase pass-through date")
        roles = Counter()
        residual = Decimal(0)
        for leg in group["legs"]:
            required_leg_fields = {"source_row", "expected_date", "expected_description",
                                   "expected_signed_amount", "expected_payment_instrument",
                                   "expected_memo", "expected_original_review_classification", "role"}
            if not isinstance(leg, dict) or set(leg) != required_leg_fields:
                raise ValueError("Invalid third-party purchase pass-through leg fields")
            row = leg["source_row"]
            if (type(row) is not int or row < 2 or row in reserved_rows
                    or row in pass_through_rows):
                raise ValueError("Duplicate or conflicting third-party purchase pass-through row")
            pass_through_rows.add(row)
            transaction = by_row.get(row)
            if (transaction is None or transaction["source_file_sha256"] != source_hash
                    or leg["expected_original_review_classification"] != "review"
                    or transaction["classification"] != leg["expected_original_review_classification"]
                    or leg["expected_date"] != group["expected_date"]
                    or any(not isinstance(leg[field], str) for field in required_leg_fields - {"source_row"})
                    or transaction["date"] != leg["expected_date"]
                    or transaction["description"] != leg["expected_description"]
                    or transaction["signed_amount"] != leg["expected_signed_amount"]
                    or transaction["payment_instrument"] != leg["expected_payment_instrument"]
                    or transaction["memo"] != leg["expected_memo"]):
                raise ValueError(f"Third-party purchase pass-through row {row} evidence mismatch")
            amount = amount_value(transaction["signed_amount"])
            role = leg["role"]
            if (amount is None or role not in {"receipt", "charge"}
                    or (role == "receipt" and amount <= 0)
                    or (role == "charge" and amount >= 0)):
                raise ValueError(f"Invalid third-party purchase pass-through role for row {row}")
            roles[role] += 1
            residual += amount
        if (roles["receipt"] < 1 or roles["charge"] != 1
                or not isinstance(group["expected_residual"], str)
                or group["expected_residual"] != amount_text(residual)):
            raise ValueError(f"Third-party purchase pass-through residual mismatch: {group_id}")
    wallet_pairs = decisions.get("confirmed_wallet_top_up_pairs", [])
    if not isinstance(wallet_pairs, list):
        raise ValueError("Invalid wallet top-up pairs")
    wallet_pair_ids = set()
    wallet_rows = set()
    reserved_rows |= pass_through_rows
    for pair in wallet_pairs:
        if (not isinstance(pair, dict)
                or set(pair) != {"group_id", "expected_date", "legs"}
                or not isinstance(pair["group_id"], str)
                or not pair["group_id"].strip()
                or pair["group_id"] != pair["group_id"].strip()
                or pair["group_id"] in wallet_pair_ids
                or not isinstance(pair["legs"], list) or len(pair["legs"]) != 2):
            raise ValueError("Invalid wallet top-up pair")
        wallet_pair_ids.add(pair["group_id"])
        try:
            pair_date = date.fromisoformat(pair["expected_date"])
        except (TypeError, ValueError):
            raise ValueError("Invalid wallet top-up date") from None
        if pair_date.isoformat() != pair["expected_date"]:
            raise ValueError("Invalid wallet top-up date")
        pair_roles = set()
        pair_total = Decimal(0)
        for leg in pair["legs"]:
            leg_fields = {"source_row", "expected_date", "expected_description",
                          "expected_signed_amount", "expected_payment_instrument",
                          "expected_memo", "expected_original_review_classification", "role"}
            if not isinstance(leg, dict) or set(leg) != leg_fields:
                raise ValueError("Invalid wallet top-up leg fields")
            row = leg["source_row"]
            if (type(row) is not int or row < 2 or row in reserved_rows or row in wallet_rows):
                raise ValueError("Duplicate or conflicting wallet top-up row")
            wallet_rows.add(row)
            transaction = by_row.get(row)
            if (transaction is None or transaction["source_file_sha256"] != source_hash
                    or leg["expected_original_review_classification"] != "review"
                    or transaction["classification"] != leg["expected_original_review_classification"]
                    or leg["expected_date"] != pair["expected_date"]
                    or any(not isinstance(leg[field], str) for field in leg_fields - {"source_row"})
                    or any(transaction[field] != leg["expected_" + field]
                           for field in ("date", "description", "signed_amount",
                                         "payment_instrument", "memo"))):
                raise ValueError(f"Wallet top-up row {row} evidence mismatch")
            amount = amount_value(transaction["signed_amount"])
            role = leg["role"]
            if (role in pair_roles or role not in {"wallet_receipt", "bank_outflow"}
                    or amount is None or (role == "wallet_receipt" and amount <= 0)
                    or (role == "bank_outflow" and amount >= 0)):
                raise ValueError(f"Invalid wallet top-up role for row {row}")
            pair_roles.add(role)
            pair_total += amount
        if pair_roles != {"wallet_receipt", "bank_outflow"} or pair_total != 0:
            raise ValueError(f"Wallet top-up pair does not balance: {pair['group_id']}")
    payment_top_ups, generic_transfers, payment_expenses = validate_payment_decisions(
        transactions, source_hash, decisions, reserved_rows | wallet_rows)
    duplicate_pairs = validate_duplicate_decisions(transactions, source_hash, decisions)
    for pair in duplicate_pairs:
        keep = by_row[pair["keep_source_row"]]
        exclude = by_row[pair["exclude_source_row"]]
        keep["aggregation_status"] = "included_owner_confirmed_duplicate"
        keep["aggregation_decision"] = {
            "paired_source_row": exclude["source_row"],
            "approval_date": pair["approval_date"],
            "treatment": pair["treatment"],
        }
        exclude["aggregation_status"] = "excluded_owner_confirmed_duplicate"
        exclude["aggregation_decision"] = {
            "paired_source_row": keep["source_row"],
            "approval_date": pair["approval_date"],
            "treatment": pair["treatment"],
        }
    for row in approved_rows:
        transaction = by_row[row]
        transaction["classification"] = "own-account transfer"
        transaction["classification_evidence"].extend([
            "owner_approved_own_account_transfer", f"owner_approval_date:{approval_date.isoformat()}",
        ])
    for confirmation in confirmations:
        transaction = by_row[confirmation["source_row"]]
        transaction["destination_account"] = confirmation["destination_account"]
        transaction["destination_evidence"] = [
            "owner_confirmed_destination",
            f"owner_confirmation_date:{confirmation['confirmation_date']}",
        ]
    for confirmation in irp_confirmations:
        transaction = by_row[confirmation["source_row"]]
        transaction["destination_account"] = confirmation["destination_account"]
        transaction["contribution_purpose"] = confirmation["contribution_purpose"]
        transaction["owner_confirmation_date"] = confirmation["confirmation_date"]
        transaction["contribution_evidence"] = [
            "owner_confirmed_irp_cash_contribution",
            f"owner_confirmation_date:{confirmation['confirmation_date']}",
        ]
    for decision in review_decisions:
        transaction = by_row[decision["source_row"]]
        transaction["classification"] = decision["classification"]
        transaction["classification_evidence"].extend([
            "owner_confirmed_review_classification",
            f"owner_approval_date:{approval_date.isoformat()}",
        ])
        if "purpose" in decision:
            transaction["purpose"] = decision["purpose"]
    for decision in new_decisions["confirmed_account_verification_credits"]:
        transaction = by_row[decision["source_row"]]
        transaction["classification"] = "account verification credit"
        transaction["purpose"] = decision["purpose"]
        transaction["classification_evidence"].extend([
            "owner_confirmed_account_verification_credit",
            f"owner_approval_date:{approval_date.isoformat()}",
        ])
    for decision in new_decisions["tentative_review_contexts"]:
        transaction = by_row[decision["source_row"]]
        transaction["owner_context"] = decision["owner_context"]
        transaction["owner_context_status"] = "tentative"
        transaction["owner_context_evidence"] = [
            "owner_tentative_context", f"owner_approval_date:{approval_date.isoformat()}",
        ]
    for decision in new_decisions["confirmed_review_purposes"]:
        transaction = by_row[decision["source_row"]]
        transaction["purpose"] = decision["purpose"]
        transaction["purpose_status"] = "confirmed_accounting_treatment_pending"
        transaction["purpose_evidence"] = [
            "owner_confirmed_purpose", f"owner_approval_date:{approval_date.isoformat()}",
        ]
    for decision in new_decisions["owner_estimated_family_transfers"]:
        transaction = by_row[decision["source_row"]]
        transaction["classification"] = "owner-estimated family transfer"
        transaction["classification_evidence"].extend([
            "owner_recollection_unverified", f"owner_approval_date:{approval_date.isoformat()}",
        ])
        transaction["owner_recollection"] = decision["owner_recollection"]
        transaction["owner_recollection_status"] = "estimate_unverified"
        transaction["candidate_related_row"] = decision["candidate_related_row"]
        transaction["relation_status"] = decision["relation_status"]
        transaction["spending_treatment"] = "excluded_non_consumption"
    for decision in new_decisions["confirmed_refund_reversals"]:
        transaction = by_row[decision["source_row"]]
        transaction["classification"] = "refund/reversal"
        transaction["classification_evidence"].extend([
            "owner_confirmed_cashback", f"owner_approval_date:{approval_date.isoformat()}",
        ])
        transaction["purpose"] = decision["purpose"]
        transaction["spending_treatment"] = decision["spending_treatment"]
    for decision in new_decisions["confirmed_consumption_expenses"]:
        transaction = by_row[decision["source_row"]]
        transaction["classification"] = "expense"
        transaction["classification_evidence"].extend([
            "owner_confirmed_consumption_expense", f"owner_approval_date:{approval_date.isoformat()}",
        ])
        transaction["spending_treatment"] = "consumption"
    for decision in new_decisions["confirmed_own_account_receipts"]:
        transaction = by_row[decision["source_row"]]
        transaction["classification"] = "own-account transfer"
        transaction["classification_evidence"].extend([
            "owner_confirmed_own_account_receipt", f"owner_approval_date:{approval_date.isoformat()}",
        ])
        transaction["income_treatment"] = "excluded_non_income"
        transaction["spending_treatment"] = "excluded_non_consumption"
    for decision in new_decisions["confirmed_group_expense_reimbursements"]:
        transaction = by_row[decision["source_row"]]
        transaction["classification"] = "expense reimbursement received"
        transaction["classification_evidence"].extend([
            "owner_confirmed_group_expense_settlement",
            f"owner_approval_date:{approval_date.isoformat()}",
        ])
        transaction["owner_context"] = decision["owner_context"]
        transaction["income_treatment"] = "excluded_non_income"
        transaction["spending_treatment"] = decision["spending_treatment"]
    for decision in new_decisions["confirmed_bank_interest_income"]:
        transaction = by_row[decision["source_row"]]
        transaction["classification"] = "income"
        transaction["income_source"] = "bank interest"
        transaction["classification_evidence"].extend([
            "owner_confirmed_bank_interest", f"owner_approval_date:{approval_date.isoformat()}",
        ])
    for decision in new_decisions["confirmed_prior_spending_reimbursements"]:
        transaction = by_row[decision["source_row"]]
        transaction["classification"] = "expense reimbursement received"
        transaction["classification_evidence"].extend([
            "owner_confirmed_prior_spending_reimbursement",
            f"owner_approval_date:{approval_date.isoformat()}",
        ])
        transaction["income_treatment"] = "excluded_non_income"
        transaction["spending_treatment"] = decision["spending_treatment"]
    for decision in new_decisions["confirmed_food_consumption_expenses"]:
        transaction = by_row[decision["source_row"]]
        transaction["classification"] = "expense"
        transaction["classification_evidence"].extend([
            "owner_confirmed_food_consumption", f"owner_approval_date:{approval_date.isoformat()}",
        ])
        transaction["spending_treatment"] = "consumption"
    for decision in new_decisions["confirmed_unmatched_wallet_top_ups"]:
        transaction = by_row[decision["source_row"]]
        transaction["classification"] = "wallet top-up"
        transaction["classification_evidence"].extend([
            "owner_confirmed_wallet_top_up_without_matched_receipt_or_purchase",
            f"owner_approval_date:{approval_date.isoformat()}",
        ])
        transaction["spending_treatment"] = "excluded_non_consumption"
    for pair in wallet_pairs:
        for leg in pair["legs"]:
            transaction = by_row[leg["source_row"]]
            transaction["classification"] = "own-account transfer"
            transaction["classification_evidence"].extend([
                "owner_confirmed_wallet_top_up_pair",
                f"owner_approval_date:{approval_date.isoformat()}",
            ])
            transaction["transfer_group_id"] = pair["group_id"]
            transaction["transfer_role"] = leg["role"]
            transaction["income_treatment"] = "excluded_non_income"
            transaction["spending_treatment"] = "excluded_non_consumption"
    for group in payment_top_ups:
        for role in ("bank_outflow", "wallet_receipt"):
            transaction = by_row[group[role]["source_row"]]
            if role == "bank_outflow":
                transaction["classification"] = "wallet top-up for purchase"
            transaction["classification_evidence"].extend([
                "owner_confirmed_top_up_with_separate_existing_purchase",
                f"owner_approval_date:{approval_date.isoformat()}",
            ])
            transaction["transfer_group_id"] = group["group_id"]
            transaction["transfer_role"] = role
            transaction["existing_purchase_row"] = group["purchase"]["source_row"]
            transaction["spending_treatment"] = "excluded_non_consumption"
            if role == "wallet_receipt":
                transaction["income_treatment"] = "excluded_non_income"
    for evidence in generic_transfers:
        transaction = by_row[evidence["source_row"]]
        transaction["classification"] = "unverified-destination transfer"
        transaction["classification_evidence"].extend([
            "owner_classified_generic_transfer", "no_wallet_usage_within_10min_in_this_workbook_only",
            "owner_2026-10-23_unverified_usage_counts_as_expense",
            f"owner_approval_date:{approval_date.isoformat()}",
        ])
        transaction["destination_status"] = "unverified"
        transaction["spending_treatment"] = "unverified_usage_expense"
        transaction["expense_label"] = "사용처 미확인 지출"
    for evidence in payment_expenses:
        transaction = by_row[evidence["source_row"]]
        transaction["classification"] = "expense"
        transaction["classification_evidence"].extend([
            "owner_confirmed_purchase_paid_by_transfer",
            f"owner_approval_date:{approval_date.isoformat()}",
        ])
        transaction["spending_treatment"] = "consumption"
    for group in pass_through_groups:
        for leg in group["legs"]:
            transaction = by_row[leg["source_row"]]
            transaction["classification"] = "third-party purchase pass-through"
            transaction["classification_evidence"].extend([
                "owner_confirmed_purchase_reimbursement",
                f"owner_approval_date:{approval_date.isoformat()}",
            ])
            transaction["pass_through_group_id"] = group["group_id"]
            transaction["pass_through_role"] = leg["role"]
            transaction["relation_status"] = group["relation_status"]
            transaction["group_residual"] = group["expected_residual"]
            transaction["residual_status"] = "unresolved_difference" if group["expected_residual"] != "0" else "zero"
            transaction["spending_treatment"] = "excluded_non_consumption"
            if leg["role"] == "receipt":
                transaction["income_treatment"] = "excluded_non_income"


def reconciliation_candidates(transactions):
    card_settlements = [item for item in transactions if item["classification"] == "card settlement"]
    settlement_candidates = []
    for settlement in card_settlements:
        card_name = next((name for name in ("삼성", "신한", "하나", "국민", "현대", "롯데", "우리")
                          if name in settlement["description"]), None)
        candidate = {"settlement_row": settlement["source_row"],
                     "card_brand": card_name, "status": "billing_cycle_unverified"}
        if card_name == "신한":
            try:
                settlement_day = date.fromisoformat(settlement["date"])
            except ValueError:
                settlement_day = None
            if settlement_day:
                previous_month_key = date.fromordinal(
                    settlement_day.replace(day=1).toordinal() - 1).strftime("%Y-%m")
                kpass_rows = [item for item in transactions
                              if item["classification"] == "expense"
                              and item["date"].startswith(previous_month_key)
                              and item["original_category"] == "교통"
                              and item["original_subcategory"] == "대중교통"
                              and "K-패스" in item["payment_instrument"]]
                kpass_total = sum((-amount_value(item["signed_amount"]) for item in kpass_rows),
                                  Decimal(0))
                if kpass_rows and kpass_total == -amount_value(settlement["signed_amount"]):
                    candidate["previous_month_kpass_candidate"] = {
                        "purchase_month": previous_month_key,
                        "purchase_rows": [item["source_row"] for item in kpass_rows],
                        "purchase_total": amount_text(kpass_total),
                        "reason": "previous_calendar_month_kpass_total_equals_shinhan_settlement",
                        "status": "candidate_only_billing_cycle_unverified",
                    }
        settlement_candidates.append(candidate)
    transit_groups = defaultdict(list)
    for item in transactions:
        if item["original_category"] == "교통" and item["original_subcategory"] == "대중교통":
            key = (item["date"], item["description"], item["signed_amount"],
                   item["payment_instrument"])
            transit_groups[key].append(item["source_row"])
    transit_candidates = [{"source_rows": rows, "status": "candidate_only_no_auto_drop"}
                          for rows in transit_groups.values() if len(rows) > 1]
    return {"card_settlements": settlement_candidates,
            "transit": sorted(transit_candidates, key=lambda group: group["source_rows"][0])}


def spending_overlap_candidates(transactions, spending_dir):
    """Use transaction-level evidence only; never remove a possible duplicate."""
    if spending_dir is None:
        return []
    matches = []
    by_date_amount = defaultdict(list)
    card_brands = ("삼성", "신한", "하나", "국민", "현대", "롯데", "우리")

    def normalize_merchant(value):
        return re.sub(r"\s+", "", str(value or "")).casefold()

    for transaction in transactions:
        if transaction["classification"] == "expense":
            by_date_amount[(transaction["date"], transaction["signed_amount"])].append(transaction)
    for path in sorted(spending_dir.rglob("*.summary.json")):
        data = json.loads(path.read_text(encoding="utf-8"))
        records = data.get("entries", []) if isinstance(data, dict) else []
        if not isinstance(records, list):
            continue
        for index, record in enumerate(records):
            if not isinstance(record, dict) or record.get("status") != "expense":
                continue
            amount = amount_value(record.get("amount"))
            merchant = normalize_merchant(record.get("merchant"))
            card_brand = str(record.get("card") or "")
            if amount is None or amount <= 0:
                continue
            record_date = str(record.get("at") or data.get("date") or "")[:10]
            candidates = by_date_amount.get((record_date, amount_text(-amount)), [])
            if not candidates:
                continue
            brand_matches = [item for item in candidates
                             if card_brand in card_brands
                             and card_brand in item["payment_instrument"]]
            merchant_matches = [item for item in brand_matches
                                if merchant and merchant == normalize_merchant(item["description"])]
            if len(merchant_matches) == 1:
                candidates, tier, reason = merchant_matches, "high", "date_amount_card_brand_merchant"
            elif len(merchant_matches) > 1:
                candidates, tier, reason = merchant_matches, None, "ambiguous_exact_merchant"
            elif len(brand_matches) == 1:
                candidates, tier, reason = brand_matches, "medium", "unique_date_amount_card_brand"
            elif len(brand_matches) > 1:
                candidates, tier, reason = brand_matches, None, "ambiguous_card_brand"
            elif len(candidates) == 1 and card_brand in card_brands:
                tier, reason = "low", "unique_date_amount_instrument_differs"
            elif len(candidates) == 1:
                tier, reason = None, "card_brand_unverified"
            else:
                tier, reason = None, "ambiguous_date_amount"
            candidate_rows = [item["source_row"] for item in candidates]
            candidate = {"summary_file": str(path.relative_to(spending_dir)),
                         "summary_record_index": index, "candidate_rows": candidate_rows,
                         "tier": tier, "reason": reason,
                         "status": "potential_overlap_no_auto_drop" if tier else "ambiguous_no_auto_drop"}
            if tier:
                candidate["source_row"] = candidate_rows[0]
            matches.append(candidate)
    return matches


def monthly_breakdown(transactions, source_coverage=None):
    """Positive monthly amounts; finalize months only when source coverage is known."""
    dated = [item["date"] for item in transactions if item["date"]]
    if not dated:
        return []
    first_date, last_date = min(dated), max(dated)
    months = {}
    amount_fields = ("total_expense", "everyday_expense", "one_off_tax",
                     "confirmed_prepayment", "unverified_usage_spending", "total_income",
                     "recurring_salary_income", "bonus_income", "other_income")
    for item in transactions:
        if not item["date"] or item["aggregation_status"] == "excluded_owner_confirmed_duplicate":
            continue
        amount = amount_value(item["signed_amount"])
        if amount is None:
            continue
        month = item["date"][:7]
        if month not in months:
            months[month] = {"amounts": {field: Decimal(0) for field in amount_fields},
                             "review_rows": [], "observed_dates": []}
        bucket = months[month]
        bucket["observed_dates"].append(item["date"])
        values = bucket["amounts"]
        if amount < 0 and (item["classification"] == "expense"
                           or item.get("spending_treatment") == "unverified_usage_expense"):
            spent = -amount
            values["total_expense"] += spent
            if item.get("spending_treatment") == "unverified_usage_expense":
                values["unverified_usage_spending"] += spent
            elif (item["original_category"] == "금융"
                  and item["original_subcategory"] == "세금/과태료"):
                values["one_off_tax"] += spent
                if item["memo"].strip() not in ONE_OFF_TAX_MEMOS:
                    bucket["review_rows"].append(item["source_row"])
            elif ((item["original_category"] == "뷰티/미용"
                   and item["memo"].strip() == "미용실 선불")
                  or (item["original_category"], item["original_subcategory"],
                      item["description"], item["memo"].strip()) in {
                          ("의료/건강", "운동", "국가대표 복싱센터 송도2호점", "복싱 6개월 주 5회"),
                          ("문화/여가", "운동", "국가대표 복싱센터", "3개월"),
                          ("온라인쇼핑", "서비스구독", "AMAZON BD", "유튜브 인도 6개월"),
                          ("온라인쇼핑", "서비스구독", "AMAZON BD", "유튜브 3개월"),
                          ("교육/학습", "미분류", "고래수학교습소", "미네 수학 7월, 8월"),
                      }):
                values["confirmed_prepayment"] += spent
            else:
                values["everyday_expense"] += spent
        elif amount > 0 and item["classification"] == "income":
            values["total_income"] += amount
            if (item["original_category"] == "급여"
                    and item["description"] == "전자급여"):
                values["recurring_salary_income"] += amount
            elif item["original_category"] == "상여금":
                values["bonus_income"] += amount
            else:
                values["other_income"] += amount
    result = []
    for month, bucket in sorted(months.items()):
        month_end = f"{month}-{calendar.monthrange(int(month[:4]), int(month[5:7]))[1]:02}"
        if source_coverage and (f"{month}-01" < source_coverage[0]
                                or month_end > source_coverage[1]):
            continue
        values = bucket["amounts"]
        result.append({
            "month": month,
            "observed_first_date": min(bucket["observed_dates"]),
            "observed_last_date": max(bucket["observed_dates"]),
            "boundary_month_partial": (False if source_coverage else (
                (month == first_date[:7] and first_date[-2:] != "01")
                or (month == last_date[:7] and last_date != month_end))),
            "review_rows": sorted(bucket["review_rows"]),
            **{field: amount_text(values[field]) for field in amount_fields},
        })
    return result


def resolve_source_coverage(workbook_path, source_hash):
    if source_hash in SOURCE_COVERAGE_BY_SHA256:
        return SOURCE_COVERAGE_BY_SHA256[source_hash]
    path = Path(workbook_path)
    if (path.name == "source.xlsx"
            and path.parent.parent.name == "Banksalad"
            and path.parent.parent.parent.name == "Spending"):
        match = EXPORT_DATE_RANGE.fullmatch(path.parent.name)
        if match:
            try:
                start, end = (date.fromisoformat(value) for value in match.groups())
            except ValueError:
                pass
            else:
                if start <= end:
                    return start.isoformat(), end.isoformat()
    raise ValueError("Source coverage unknown: expected approved SHA or Banksalad export date range")


def build_report(workbook_path, spending_dir=None, owner_decisions=None):
    source_hash, transactions = load_transactions(workbook_path)
    source_coverage = resolve_source_coverage(workbook_path, source_hash)
    mark_paired_own_account_transfers(transactions)
    pass_through_groups = []
    if owner_decisions is not None:
        decisions = json.loads(owner_decisions.read_text(encoding="utf-8"))
        apply_owner_decisions(transactions, source_hash, decisions)
        pass_through_groups = [
            {"group_id": group["group_id"], "date": group["expected_date"],
             "source_rows": [leg["source_row"] for leg in group["legs"]],
             "signed_residual": group["expected_residual"],
             "residual_status": "unresolved_difference" if group["expected_residual"] != "0" else "zero",
             "relation_status": group["relation_status"]}
            for group in decisions.get("confirmed_third_party_purchase_pass_throughs", [])
        ]
    duplicates = mark_exact_duplicates(transactions)
    category_totals = defaultdict(lambda: Decimal(0))
    classification_totals = defaultdict(lambda: Decimal(0))
    raw_category_totals = defaultdict(lambda: Decimal(0))
    raw_classification_totals = defaultdict(lambda: Decimal(0))
    counts = Counter()
    for item in transactions:
        amount = amount_value(item["signed_amount"])
        counts[item["classification"]] += 1
        if amount is not None:
            category_key = (item["original_type"], item["original_category"],
                            item["original_subcategory"])
            raw_category_totals[category_key] += amount
            raw_classification_totals[item["classification"]] += amount
            if item["aggregation_status"] != "excluded_owner_confirmed_duplicate":
                category_totals[category_key] += amount
                classification_totals[item["classification"]] += amount
    return {
        "schema_version": 1,
        "source_file_sha256": source_hash,
        "source_sheet": SHEET_NAME,
        "coverage": {
            "transaction_rows": len(transactions),
            "classified_rows": len(transactions) - counts["review"],
            "review_rows": counts["review"],
            "classification_counts": {name: counts[name] for name in CLASSIFICATIONS},
            "effective_aggregation_rows": sum(item["aggregation_status"] !=
                                              "excluded_owner_confirmed_duplicate"
                                              for item in transactions),
            "excluded_approved_duplicate_rows": sum(item["aggregation_status"] ==
                                                     "excluded_owner_confirmed_duplicate"
                                                     for item in transactions),
            "raw_classification_signed_totals": {name: amount_text(raw_classification_totals[name])
                                                 for name in CLASSIFICATIONS},
            "classification_signed_totals": {name: amount_text(raw_classification_totals[name])
                                             for name in CLASSIFICATIONS},
            "effective_classification_signed_totals": {name: amount_text(classification_totals[name])
                                                       for name in CLASSIFICATIONS},
            "exact_duplicate_groups": len(duplicates),
            "exact_duplicate_rows": sum(group["count"] for group in duplicates),
            "missing_date_rows": sum(not item["date"] for item in transactions),
            "missing_instrument_rows": sum(not item["payment_instrument"] for item in transactions),
        },
        "effective_category_totals": [
            {"original_type": kind, "original_category": category,
             "original_subcategory": subcategory, "signed_total": amount_text(total)}
            for (kind, category, subcategory), total in sorted(category_totals.items())
        ],
        "raw_category_totals": [
            {"original_type": kind, "original_category": category,
             "original_subcategory": subcategory, "signed_total": amount_text(total)}
            for (kind, category, subcategory), total in sorted(raw_category_totals.items())
        ],
        "category_totals": [
            {"original_type": kind, "original_category": category,
             "original_subcategory": subcategory, "signed_total": amount_text(total)}
            for (kind, category, subcategory), total in sorted(raw_category_totals.items())
        ],
        "monthly_breakdown": monthly_breakdown(
            transactions, source_coverage),
        "review_queue": [{"source_row": item["source_row"],
                          "evidence": item["classification_evidence"]}
                         for item in transactions if item["classification"] == "review"],
        "exact_duplicate_groups": duplicates,
        "reconciliation_candidates": reconciliation_candidates(transactions),
        "spending_overlap_candidates": spending_overlap_candidates(transactions, spending_dir),
        "third_party_purchase_pass_throughs": {
            "groups": pass_through_groups,
            "aggregate_signed_residual": amount_text(sum(
                (Decimal(group["signed_residual"]) for group in pass_through_groups), Decimal(0))),
        },
        "transactions": transactions,
    }


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", required=True, type=Path,
                        help="Explicit JSON report destination; no default output")
    parser.add_argument("--spending-dir", type=Path,
                        help="Read existing transaction-level summary JSONs for tiered overlap candidates")
    parser.add_argument("--owner-decisions", type=Path,
                        help="JSON file with source hash, approved own-transfer rows, and optional destinations")
    arguments = parser.parse_args()
    if arguments.output.resolve() == arguments.workbook.resolve():
        parser.error("Output may not overwrite the source workbook")
    if arguments.spending_dir and not arguments.spending_dir.is_dir():
        parser.error("--spending-dir must be an existing directory")
    if arguments.spending_dir and arguments.output.resolve().is_relative_to(arguments.spending_dir.resolve()):
        parser.error("Output may not be inside --spending-dir")
    if arguments.owner_decisions and arguments.output.resolve() == arguments.owner_decisions.resolve():
        parser.error("Output may not overwrite --owner-decisions")
    try:
        report = build_report(arguments.workbook, arguments.spending_dir, arguments.owner_decisions)
    except (OSError, ValueError, json.JSONDecodeError) as error:
        parser.error(str(error))
    arguments.output.parent.mkdir(parents=True, exist_ok=True)
    arguments.output.write_text(json.dumps(report, ensure_ascii=False, sort_keys=True,
                                           indent=2) + "\n", encoding="utf-8")
    print("Audit report written; transaction count:", report["coverage"]["transaction_rows"])


if __name__ == "__main__":
    main()
